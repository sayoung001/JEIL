"""월간 실적보고서 자동 집계 (진단 PART D §10).

진단이 **실제 8월 보고서와 대조한 표**가 정답지다.

    동재하 실시/합격      3 / 3   <- O-01,04 재하시험 시트 3행
    겉모양 및 모양        4 / 4   <- Q-01,02 26.08 의 Q-Q-02-NN
    치수                  4 / 4   <- 같음 (시험 1건이 보고서 2줄)
    용접부 비파괴검사     1 / 1   <- O-01,04 MT검사 시트 1행
"""
from __future__ import annotations

from datetime import date

import openpyxl
import pytest

from core.config import Config
from core.ledger import COL, FIRST_DATA_ROW
from core.monthly_report import MonthlyReporter, load_매핑
from core.util import sha256
from tests.fixtures import cfg as base_cfg


def _헤더(ws) -> None:
    ws.cell(3, COL["일련번호"], "일련\n번호")
    ws.cell(3, COL["날짜"], "날 짜")
    ws.cell(3, COL["구분"], "시험·검사\n구분")


def _겉모양(ws, row: int, 번호: int, 연번: int, 날짜: date, 판정: str = "합 격") -> int:
    """겉모양·치수 1건 = 5행 그룹. A열은 첫 행에만."""
    ws.cell(row, COL["일련번호"], 연번)
    ws.cell(row, COL["날짜"], 날짜)
    ws.cell(row, COL["구분"], f"Q-Q-\n02-{번호:02d}")
    ws.cell(row, COL["대상재료"], "PHC 파일 \n겉모양, 치수")
    ws.cell(row, COL["종목"], "치 수")
    ws.cell(row, COL["판정"], 판정)
    ws.cell(row + 3, COL["종목"], "모양")
    ws.cell(row + 4, COL["종목"], "겉모양")
    return row + 5


def _밀크(ws, row: int, 번호: int, 연번: int, 날짜: date, 판정: str = "합 격") -> int:
    ws.cell(row, COL["일련번호"], 연번)
    ws.cell(row, COL["날짜"], 날짜)
    ws.cell(row, COL["구분"], f"Q-Q-\n01-{번호:02d}")
    ws.cell(row, COL["대상재료"], "PHC파일 밀크")
    ws.cell(row, COL["판정"], 판정)
    return row + 2


def _재하MT(ws, row: int, 연번: int, 날짜: date, 판정: str = "합 격",
            비고: str = "") -> int:
    ws.cell(row, COL["일련번호"], 연번)
    ws.cell(row, COL["날짜"], 날짜)
    ws.cell(row, COL["구분"], "의뢰시험")
    ws.cell(row, COL["대상재료"], "PHC 파일")
    ws.cell(row, COL["판정"], 판정)
    if 비고:
        ws.cell(row, COL["비고"], 비고)
    return row + 1


@pytest.fixture()
def 대장(tmp_path):
    """진단 §10 의 8월 실측 상황을 그대로 만든다."""
    def build(*, 겉모양=4, 재하=3, mt=1, 밀크=0, 칠월=0, 판정="합 격"):
        # --- Q-01,02 (월 시트형) ---
        pile = tmp_path / "Q-01,02.xlsx"
        wb = openpyxl.Workbook()
        for 달, 건수 in (("26.07", 칠월), ("26.08", 겉모양)):
            ws = wb.create_sheet(달)
            _헤더(ws)
            row, 연번 = FIRST_DATA_ROW, 1
            for i in range(건수):
                row = _겉모양(ws, row, i + 1, 연번, date(2026, int(달[3:]), 5 + i), 판정)
                연번 += 1
            for i in range(밀크 if 달 == "26.08" else 0):
                row = _밀크(ws, row, i + 1, 연번, date(2026, 8, 10 + i))
                연번 += 1
        del wb["Sheet"]
        wb.save(pile)

        # --- Q-03 (월 시트형, 비어 있음) ---
        civil = tmp_path / "Q-03.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "26.08"
        _헤더(ws)
        wb.save(civil)

        # --- N-02 (월 시트형, 비어 있음) ---
        outsrc = tmp_path / "N-02.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "26.08"
        _헤더(ws)
        wb.save(outsrc)

        # --- O-01,04 (종류 시트형) ---
        load_mt = tmp_path / "O-01,04.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "재하시험"
        _헤더(ws)
        row, 연번 = FIRST_DATA_ROW, 1
        for i in range(재하):
            row = _재하MT(ws, row, 연번, date(2026, 8, 24 + i), 판정)
            연번 += 1
        row = _재하MT(ws, row, 연번, date(2026, 7, 30))      # 7월 건 — 8월에 세면 안 된다
        ws2 = wb.create_sheet("MT검사")
        _헤더(ws2)
        row, 연번 = FIRST_DATA_ROW, 1
        for i in range(mt):
            row = _재하MT(ws2, row, 연번, date(2026, 8, 24 + i), 판정)
            연번 += 1
        wb.save(load_mt)

        c = base_cfg()
        raw = dict(c.raw)
        raw["files"] = dict(raw["files"])
        raw["files"].update(ledger_pile=str(pile), ledger_civil=str(civil),
                            ledger_outsrc=str(outsrc), ledger_load_mt=str(load_mt))
        return Config(raw=raw)
    return build


def _reporter(cfg) -> MonthlyReporter:
    return MonthlyReporter(cfg)


# =====================================================================
# 진단 §10 실측표 재현
# =====================================================================
def test_8월_실측표를_그대로_재현한다(대장):
    """이 테스트가 이 기능의 존재 이유다."""
    t = _reporter(대장()).tally("26.08")
    종목 = t.종목

    assert (종목["동재하시험"].실시, 종목["동재하시험"].합격) == (3, 3)
    assert (종목["겉모양 및 모양"].실시, 종목["겉모양 및 모양"].합격) == (4, 4)
    assert (종목["치수"].실시, 종목["치수"].합격) == (4, 4)
    assert (종목["용접부 비파괴검사"].실시, 종목["용접부 비파괴검사"].합격) == (1, 1)
    assert t.경고 == []


def test_겉모양_시험1건이_보고서_두줄로_잡힌다(대장):
    """4번 시험했는데 보고서에는 겉모양 4 · 치수 4 로 각각 잡힌다."""
    t = _reporter(대장(겉모양=4)).tally("26.08")
    assert t.종목["겉모양 및 모양"].실시 == 4
    assert t.종목["치수"].실시 == 4
    # 실제 시험 건수는 4건 (같은 출처를 두 종목이 공유한다)
    assert set(t.종목["겉모양 및 모양"].출처) == set(t.종목["치수"].출처)
    assert len(t.종목["치수"].출처) == 4


def test_행이_아니라_행그룹을_센다(대장):
    """겉모양 1건은 5행이다. 행을 세면 5배로 부푼다."""
    t = _reporter(대장(겉모양=4)).tally("26.08")
    assert t.종목["치수"].실시 == 4          # 20 이 아니다


# =====================================================================
# 달 가르기
# =====================================================================
def test_종류시트형은_날짜로_달을_가른다(대장):
    """O-01,04 는 시트가 달이 아니다. 7월 건이 8월에 섞이면 안 된다."""
    cfg = 대장(재하=3)
    팔월 = _reporter(cfg).tally("26.08")
    칠월 = _reporter(cfg).tally("26.07")
    assert 팔월.종목["동재하시험"].실시 == 3
    assert 칠월.종목["동재하시험"].실시 == 1


def test_월시트형은_해당_시트만_본다(대장):
    cfg = 대장(겉모양=4, 칠월=2)
    assert _reporter(cfg).tally("26.08").종목["치수"].실시 == 4
    assert _reporter(cfg).tally("26.07").종목["치수"].실시 == 2


def test_없는_달은_빈_집계와_안내(대장):
    t = _reporter(대장()).tally("26.12")
    assert t.총건수 == 0
    assert any("26.12 시트가 없습니다" in w for w in t.경고)


# =====================================================================
# 판정
# =====================================================================
def test_불합격이_따로_세어진다(대장):
    t = _reporter(대장(겉모양=2, 판정="불합격")).tally("26.08")
    assert t.종목["치수"].실시 == 2
    assert t.종목["치수"].합격 == 0
    assert t.종목["치수"].불합격 == 2


def test_판정이_비면_경고한다(대장):
    t = _reporter(대장(겉모양=2, 판정="")).tally("26.08")
    assert t.종목["치수"].미판정 == 2
    assert any("판정이 비어 있는" in w for w in t.경고)


def test_재시험은_비고로_잡는다(tmp_path):
    from core.monthly_report import LedgerEntry

    e = LedgerEntry(대장="O-01,04", 시트="재하시험", 행=4, 일련번호=1,
                    날짜=date(2026, 8, 24), 구분="의뢰시험", 대상재료="PHC 파일",
                    종목="", 판정="합 격", 비고="재시험")
    assert e.재시험 is True
    assert e.합격 is True


# =====================================================================
# 매핑에 없는 시험 — 조용히 빠지면 안 된다
# =====================================================================
def test_모르는_종목은_조용히_빼지_않고_경고한다(tmp_path, 대장):
    cfg = 대장(겉모양=1)
    import openpyxl as ox

    wb = ox.load_workbook(cfg.path("ledger_civil"))
    ws = wb["26.08"]
    ws.cell(FIRST_DATA_ROW, COL["일련번호"], 1)
    ws.cell(FIRST_DATA_ROW, COL["날짜"], date(2026, 8, 20))
    ws.cell(FIRST_DATA_ROW, COL["구분"], "Q-Q-\n09-01")      # 매핑에 없는 계열
    ws.cell(FIRST_DATA_ROW, COL["대상재료"], "새로운시험")
    ws.cell(FIRST_DATA_ROW, COL["판정"], "합 격")
    wb.save(cfg.path("ledger_civil"))

    t = _reporter(cfg).tally("26.08")
    assert any("세지 못했습니다" in w and "종목매핑" in w for w in t.경고)


# =====================================================================
# 누계 · 분기
# =====================================================================
def test_누계는_대장에서_직접_센다(대장):
    """지난 보고서를 이어받지 않는다. 같은 달 보고서가 3본이라 정본을 모른다."""
    cfg = 대장(겉모양=4, 칠월=2)
    전월, 금월 = _reporter(cfg).누계("26.08")
    assert 전월.종목["치수"].실시 == 2
    assert 금월.종목["치수"].실시 == 4
    누계 = 전월.더하기(금월)
    assert 누계.종목["치수"].실시 == 6


def test_첫달이면_누계가_금월과_같다(대장):
    전월, 금월 = _reporter(대장(겉모양=4, 칠월=0)).누계("26.07")
    assert 전월.총건수 == 0
    assert any("이전 달 기록이" in w for w in 전월.경고)


def test_분기_누계(대장):
    cfg = 대장(겉모양=4, 칠월=2)
    q3 = _reporter(cfg).분기(2026, 3)       # 7·8·9월
    assert q3.종목["치수"].실시 == 6


def test_있는_달_목록(대장):
    assert _reporter(대장(칠월=2)).있는_달() == ["26.07", "26.08"]


# =====================================================================
# 읽기 전용
# =====================================================================
def test_집계는_대장을_수정하지_않는다(대장):
    cfg = 대장()
    keys = ("ledger_pile", "ledger_civil", "ledger_outsrc", "ledger_load_mt")
    before = {k: sha256(cfg.path(k)) for k in keys}
    _reporter(cfg).누계("26.08")
    assert {k: sha256(cfg.path(k)) for k in keys} == before


def test_대장을_못읽어도_나머지는_센다(대장, tmp_path):
    cfg = 대장()
    cfg.raw["files"]["ledger_load_mt"] = str(tmp_path / "없음.xls")
    t = _reporter(cfg).tally("26.08")
    assert t.종목["치수"].실시 == 4                    # 읽은 것은 센다
    assert "동재하시험" not in t.종목
    assert any("읽지 못했습니다" in w for w in t.경고)  # 조용히 넘어가지 않는다


def test_표_출력(대장):
    t = _reporter(대장()).tally("26.08")
    text = str(t)
    assert "동재하시험" in text and "실시" in text


def test_매핑규칙이_전부_종목을_갖는다():
    for r in load_매핑():
        assert r.종목, r
        assert any((r.계열, r.시트, r.대상재료)), r
